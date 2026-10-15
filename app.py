import difflib
import io
import logging
import re
import threading
import time
from datetime import datetime
from types import SimpleNamespace

import gspread
import pandas as pd
import streamlit as st
from google.oauth2.service_account import Credentials
from gspread.utils import rowcol_to_a1

log = logging.getLogger("bank_app")

st.set_page_config(page_title="Bank Product Database", page_icon="🏦", layout="wide")

CORE = ["Bank Category", "Bank", "Account Type", "Product Name"]
META = ["Added By", "Added At"]
FIXED = CORE + META
SCOPES = [
    "https://www.googleapis.com/auth/spreadsheets",
    "https://www.googleapis.com/auth/drive",
]
NEW = "➕ Add new..."


# ---------------------------------------------------------------- access gate
def gate():
    """Optional shared passcode (set app_password in secrets). Streamlit Cloud's
    private-app invite list is the main protection; this is a second layer."""
    pw = st.secrets.get("app_password", "")
    if not pw or st.session_state.get("ok"):
        return
    st.title("🏦 Bank Product Database")
    entered = st.text_input("Team passcode", type="password")
    if entered:
        if entered == pw:
            st.session_state["ok"] = True
            st.rerun()
        else:
            st.error("Wrong passcode")
    st.stop()


# ---------------------------------------------------------------- sheet access
@st.cache_resource
def get_worksheet():
    creds = Credentials.from_service_account_info(
        dict(st.secrets["gcp_service_account"]), scopes=SCOPES
    )
    client = gspread.authorize(creds)
    sh = client.open_by_key(st.secrets["sheet_id"])
    try:
        ws = sh.worksheet("Products")
    except gspread.WorksheetNotFound:
        ws = sh.add_worksheet("Products", rows=1000, cols=60)
    if not ws.row_values(1):
        ws.update(range_name="A1", values=[FIXED])
    return ws


@st.cache_resource(ttl=20, show_spinner=False)
def load_data():
    """Read the sheet once and work out everything the screens need (recent values per feature,
    who introduced each feature, coverage). Cached for 20 s and shared by all users, so clicking
    around never repeats this work. The returned object is read-only: never modify it."""
    ws = get_worksheet()
    values = ws.get_all_values()
    header = list(values[0]) if values else list(FIXED)
    width = len(header)
    rows = [(list(r) + [""] * (width - len(r)))[:width] for r in values[1:]]
    df = pd.DataFrame(rows, columns=header, dtype=object)
    for c in FIXED:
        if c not in df.columns:
            df[c] = ""
    feats = [c for c in df.columns if c not in FIXED]

    added_by = df["Added By"].tolist()
    recent, origins, coverage = {}, {}, {}
    for f in feats:
        col = [str(v).strip() for v in df[f].tolist()]
        filled = [(i, v) for i, v in enumerate(col) if v]
        coverage[f] = len(filled)
        origins[f] = added_by[filled[0][0]] if filled else ""
        recent[f] = list(dict.fromkeys(v for _, v in reversed(filled)))[:40]  # most recent first
    top = sorted(feats, key=lambda f: (-coverage[f], f.lower()))[:15]
    last = df.iloc[-1] if len(df) else None
    sig = (len(df), tuple(feats), "" if last is None else last["Added At"], "" if last is None else last["Product Name"])
    return SimpleNamespace(df=df, feats=feats, recent=recent, origins=origins,
                           coverage=coverage, top=top, sig=sig)


@st.cache_resource
def save_lock():
    """One lock shared by every user session of this app (Streamlit runs them in one process)."""
    return threading.Lock()


def _row_number(res):
    """Sheet row the product landed in, from Google's append response (None if unavailable)."""
    try:
        rng = res["updates"]["updatedRange"]  # e.g. "Products!A42:K42"
        return int(re.search(r"[A-Z]+(\d+)", rng.split("!")[-1]).group(1))
    except Exception:
        return None


RETRY_CODES = (429, 500, 502, 503, 504)  # temporary Google problems: worth waiting and trying again


def _fresh(ws):
    """Re-read the tab's real size from Google. The cached copy can be out of date, for example
    after someone resized or edited the sheet by hand."""
    try:
        return ws.spreadsheet.worksheet(ws.title)
    except Exception:
        return ws


def explain_error(e):
    """(status code, Google's own message, plain-English hint) for a failed Google call."""
    code = getattr(e, "code", None)
    try:
        msg = str(e.error.get("message", "")) if hasattr(e, "error") else str(e)
    except Exception:
        msg = str(e)
    low = msg.lower()
    if code == 429 or "quota" in low:
        hint = "Google is limiting how fast the app can talk to the sheet. Wait about a minute and click Save again."
    elif code == 403 and ("not been used" in low or "disabled" in low):
        hint = "The Google Sheets or Drive API is switched off for the Google Cloud project. Enable both APIs."
    elif code == 403:
        hint = ("Google refused access. Check the sheet is still shared with the service-account email as "
                "Editor, and that the sheet or its owner's Drive has not been changed.")
    elif code == 404:
        hint = "Google cannot find the sheet. Check sheet_id in the app Secrets and that the sheet was not deleted."
    elif code == 400 and ("grid" in low or "range" in low):
        hint = "The sheet's size limits were hit. Someone may have changed rows or columns by hand."
    elif code and code >= 500:
        hint = "Google had a temporary problem. Wait a minute and click Save again."
    else:
        hint = "Google rejected the save. The details below say why."
    return code, msg, hint


def save_row(row: dict):
    """Append one product and return the sheet row number it landed in.
    New features become new columns at the far right. Saves are serialised with a lock and the
    header is re-read inside it, so simultaneous saves from different windows cannot clash.
    Temporary Google errors (rate limit / server hiccup) are retried with growing pauses; any other
    error gets one retry on a freshly loaded connection to the sheet, then is reported."""
    waits = [2, 5, 10]
    attempt, refreshed = 0, False
    while True:
        ws = get_worksheet()
        try:
            with save_lock():
                header = ws.row_values(1)  # read INSIDE the lock: always the latest columns
                lower = {h.lower(): h for h in header}
                row = {lower.get(k.lower(), k): v for k, v in row.items()}  # reuse existing spelling
                new_cols = [c for c in row if c not in header]
                if new_cols:
                    ws = _fresh(ws)  # use the live column count, not a remembered one
                    start = len(header) + 1
                    if start + len(new_cols) - 1 > ws.col_count:
                        ws.add_cols(len(new_cols) + 10)
                    # write only the NEW header cells, never the whole header row
                    ws.update(range_name=rowcol_to_a1(1, start), values=[new_cols])
                    header = header + new_cols
                res = ws.append_row(
                    [row.get(h, "") for h in header],
                    value_input_option="USER_ENTERED",
                    table_range="A1",
                )
            load_data.clear()
            return _row_number(res)
        except gspread.exceptions.APIError as e:
            code = getattr(e, "code", 0)
            log.warning("Save attempt failed: %s", e)
            if code in RETRY_CODES and attempt < len(waits):
                time.sleep(waits[attempt])  # wait outside the lock so others can save meanwhile
                attempt += 1
                continue
            if code not in RETRY_CODES and not refreshed:
                refreshed = True
                get_worksheet.clear()  # rebuild the connection and sheet info, then try once more
                continue
            raise


def to_excel(df: pd.DataFrame) -> bytes:
    """Fast Excel export (streams rows instead of going through pandas styling)."""
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    feats = [c for c in df.columns if c not in FIXED]
    cols = CORE + feats + META
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Products")
    ws.freeze_panes = "E2"
    for i, c in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(14, min(32, len(c) + 4))
    ws.append(cols)
    for row in df[cols].itertuples(index=False, name=None):
        ws.append([v if v != "" else None for v in row])  # skip empty cells: smaller and faster
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- helpers
def recent_first(series: pd.Series) -> list:
    """Unique non-empty values, most recently added first."""
    vals = [v for v in series.tolist() if str(v).strip()]
    return list(dict.fromkeys(reversed(vals)))


def norm(text: str) -> str:
    """'Debit_Card-Fee' and 'debit card fee' normalise to the same string."""
    return re.sub(r"[^a-z0-9]+", " ", str(text).lower()).strip()


RESERVED = {norm(x) for x in FIXED}


def feature_matches(query: str, feature: str) -> bool:
    """Search box: matches parts of a name, ignoring case, spaces, _ and -."""
    q = query.strip()
    if not q:
        return True
    squash = lambda t: re.sub(r"[^a-z0-9]", "", t.lower())
    if squash(q) and squash(q) in squash(feature):
        return True
    return all(tok in norm(feature) for tok in norm(q).split())


def similar_features(name: str, features: list):
    """Returns (similar_list, exact_match). exact_match is an existing feature that is
    the same name apart from case, spaces and _ / - ."""
    n = norm(name)
    if not n:
        return [], None
    exact = next((f for f in features if norm(f) == n), None)
    if exact:
        return [], exact
    n_tokens = set(n.split())
    # how many existing features use each word; very common words (fee, rate...) say little
    word_freq = {}
    for f in features:
        for t in set(norm(f).split()):
            word_freq[t] = word_freq.get(t, 0) + 1
    scored = []
    for f in features:
        fn = norm(f)
        f_tokens = set(fn.split())
        ratio = difflib.SequenceMatcher(None, n, fn).ratio()
        overlap = len(n_tokens & f_tokens) / max(1, min(len(n_tokens), len(f_tokens)))
        contains = min(len(n), len(fn)) >= 4 and (n in fn or fn in n)
        rare_shared = any(len(t) >= 3 and not t.isdigit() and word_freq.get(t, 0) <= 2
                          for t in n_tokens & f_tokens)
        score = max(ratio, 0.9 * overlap if overlap >= 0.5 else 0, 0.85 if contains else 0,
                    0.65 if rare_shared else 0)
        if score >= 0.6:
            scored.append((score, f))
    scored.sort(key=lambda t: (-t[0], t[1].lower()))
    return [f for _, f in scored[:5]], None


def _set_search(feature: str, version: int):
    """Button callback: jump the feature search box to a given feature."""
    st.session_state[f"fq_{version}"] = feature


def feature_cols(df):
    return [c for c in df.columns if c not in FIXED]


def pick(label, options, key, default=None):
    """Selectbox of recent values + 'Add new' option that reveals a text box."""
    opts = options + [NEW]
    idx = opts.index(default) if default in opts else 0
    v = st.session_state.get("form_ver", 0)
    choice = st.selectbox(label, opts, index=idx, key=f"{key}_sel_{v}")
    if choice == NEW:
        return st.text_input(f"New {label.lower()}", key=f"{key}_new_{v}").strip()
    return choice


# ---------------------------------------------------------------- app
gate()
st.title("🏦 Bank Product Database")

try:
    data = load_data()
except Exception as e:
    st.error(
        "Could not connect to Google Sheets. Check `sheet_id`, the service-account secrets, "
        "and that the sheet is shared with the service-account email as Editor."
    )
    st.exception(e)
    st.stop()

df, feats, origins = data.df, data.feats, data.origins
st.session_state["seen_sig"] = data.sig  # what this page was built from (see the freshness banner)
PAGE = 30  # feature fields drawn at a time: keeps the form fast with hundreds of features


# ---------------------------------------------------------------- who is working?
def known_members(frame):
    names = list(st.secrets.get("members", [])) + recent_first(frame["Added By"])
    seen, out = set(), []
    for n in names:
        n = str(n).strip()
        if n and n.lower() not in seen:
            seen.add(n.lower())
            out.append(n)
    return sorted(out, key=str.lower)


with st.sidebar:
    st.header("👤 Your workspace")
    members = known_members(df)
    qp_name = str(st.query_params.get("member", "")).strip()
    if qp_name and qp_name.lower() not in [m.lower() for m in members]:
        members = sorted(members + [qp_name], key=str.lower)
    start = next((i for i, m in enumerate(members) if m.lower() == qp_name.lower()), None) if qp_name else None
    me = st.selectbox(
        "Who are you?", members, index=start, accept_new_options=True,
        placeholder="Pick or type your name", key="member_sel",
    )
    me = (me or "").strip()
    if me and me != qp_name:
        st.query_params["member"] = me  # makes a personal link: ...app-url/?member=YourName
    if me:
        st.success(f"Working as **{me}**")
        st.caption(
            "Bookmark this page: the link remembers your name. Everyone works in their own window "
            "but shares one feature list and one Google Sheet."
        )
    else:
        st.info("Pick your name to start (type it if you're new).")

mine = df[df["Added By"].str.lower() == me.lower()] if me else df.iloc[0:0]
latest = feats[-5:][::-1]  # right-most columns are the newest features
new_from_others = [f for f in latest if me and origins.get(f, "").strip().lower() not in ("", me.lower())]

c1, c2, c3, c4, c5 = st.columns(5)
c1.metric("Team products", len(df))
c2.metric("My products", len(mine) if me else "-")
c3.metric("Banks", df["Bank"].replace("", pd.NA).nunique())
c4.metric("Features tracked", len(feats))
c5.metric("Account types", df["Account Type"].replace("", pd.NA).nunique())


def freshness_banner():
    """Called each time the form is used: if teammates added data since this page loaded,
    offer a one-click reload. (No background timer: it only checks when you interact.)"""
    try:
        latest_data = load_data()  # cached for 20 s and shared by everyone, so this is almost free
    except Exception:
        return
    if latest_data.sig != st.session_state.get("seen_sig"):
        a, b = st.columns([5, 1])
        a.info("🔄 Teammates added new products or features since this page loaded.")
        if b.button("Load latest", key="load_latest"):
            st.rerun()


# ---------------------------------------------------------------- shared widgets
def feature_column_picker(key: str):
    """Which feature columns a wide table shows. Showing hundreds at once is slow and unreadable;
    the Excel download always contains everything."""
    if st.checkbox("Show all feature columns (slower)", key=f"allcols_{key}"):
        return feats
    return st.multiselect(
        f"Feature columns to show ({len(feats)} exist; Excel has all of them)",
        feats, default=data.top, key=f"cols_{key}",
    )


def excel_download(frame: pd.DataFrame, label: str, prefix: str, key: str):
    """Build the Excel file only when asked (it takes a couple of seconds), then offer the download."""
    if st.button(f"📦 Prepare {label}", key=f"prep_{key}"):
        with st.spinner("Building the Excel file..."):
            st.session_state[f"xl_{key}"] = (to_excel(frame), datetime.now())
    blob = st.session_state.get(f"xl_{key}")
    if blob:
        st.download_button(
            f"⬇️ Download {label} (prepared {blob[1]:%H:%M})",
            data=blob[0],
            file_name=f"{prefix}_{blob[1]:%Y%m%d_%H%M}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            key=f"dl_{key}",
            on_click="ignore",
            type="primary",
        )


def rerun_form():
    """Redraw just the form (fast). If Streamlit says a form-only redraw is not possible right now
    (e.g. a full-page refresh was already in progress), fall back to redrawing the whole page."""
    try:
        st.rerun(scope="fragment")
    except st.errors.StreamlitAPIException:
        st.rerun()


def _goto(key, label):
    st.session_state[key] = label


tab_add, tab_mine, tab_browse, tab_overview = st.tabs(
    ["➕ Add product", "👤 My work", "🔎 All data & download", "📊 Team overview"]
)


# ------------------------------------------------------------------ ADD
@st.fragment
def add_form():
    # A fragment: typing or picking here only redraws this form, not the whole page.
    ver = st.session_state.setdefault("form_ver", 0)
    if "flash" in st.session_state:
        st.success(st.session_state.pop("flash"))
    freshness_banner()
    left, right = st.columns([1, 1.4], gap="large")

    with left:
        st.subheader("Product details")
        if me:
            st.caption(f"Adding as **{me}**. Your banks are listed first.")
        else:
            st.warning("Pick your name in the sidebar before saving.")

        my_banks = recent_first(mine["Bank"])
        banks = my_banks + [b for b in recent_first(df["Bank"]) if b not in my_banks]
        bank = pick("Bank", banks, "bank")

        # category auto-fills from the bank's earlier entries
        known_cat = None
        if bank and bank in set(df["Bank"]):
            known_cat = df.loc[df["Bank"] == bank, "Bank Category"].iloc[-1] or None
        categories = recent_first(df["Bank Category"])
        category = pick("Bank Category", categories, f"cat_{bank}", default=known_cat)

        acc_types = recent_first(df["Account Type"])
        acc_type = pick("Account Type", acc_types, "acc")

        product = st.text_input("Product Name", key=f"product_{ver}").strip()

        if bank and product:
            dup = df[(df["Bank"] == bank) & (df["Product Name"].str.lower() == product.lower())]
            if len(dup):
                st.warning("This bank already has a product with this name.")

        if bank and bank in set(df["Bank"]):
            st.caption(f"{bank} already has {int((df['Bank'] == bank).sum())} product(s) in the database.")

    with right:
        st.subheader("Features")

        # Remember typed values even while a feature is off-screen (hidden by the search box or on
        # another page). Streamlit forgets widgets that are not drawn, so we copy values aside.
        fvals = st.session_state.setdefault("fvals", {})
        suffix = f"_{ver}"
        for k in list(st.session_state.keys()):
            if k.startswith("fv_") and k.endswith(suffix):
                fname = k[3:-len(suffix)]
                cur = st.session_state[k]
                if cur is not None and str(cur).strip():
                    fvals[fname] = str(cur).strip()
                else:
                    fvals.pop(fname, None)

        if latest:
            with st.expander("🆕 Latest features added by the team"):
                for f in latest:
                    st.write(f"• **{f}**, first used by {origins.get(f) or 'unknown'}")

        s1, s2 = st.columns([3, 2])
        q = s1.text_input("🔍 Find a feature", key=f"fq_{ver}",
                          placeholder="type part of a name, e.g. card, atm, insurance")
        filled_only = s2.checkbox("Only ones I've filled", key=f"fo_{ver}")

        pool = [f for f in sorted(feats, key=str.lower) if feature_matches(q, f)]  # A-Z
        if filled_only:
            pool = [f for f in pool if f in fvals]
        n_pages = max(1, -(-len(pool) // PAGE))
        pg = 0
        if n_pages > 1:
            labels = [f"{i + 1}/{n_pages}: {pool[i * PAGE][:20]} … {pool[min((i + 1) * PAGE, len(pool)) - 1][:20]}"
                      for i in range(n_pages)]
            pkey = f"pg_{ver}_{abs(hash(tuple(labels))) % 10**8}"  # new key when the list changes
            if pkey not in st.session_state or st.session_state[pkey] not in labels:
                st.session_state[pkey] = labels[0]
            pg = labels.index(st.session_state[pkey])
            b1, b2, b3 = st.columns([1, 4, 1])
            b1.button("◀ Previous", key=f"prev_{pkey}", disabled=pg == 0, on_click=_goto,
                      args=(pkey, labels[max(pg - 1, 0)]))
            b2.selectbox("Jump to page (features are A-Z)", labels, key=pkey)
            b3.button("Next ▶", key=f"next_{pkey}", disabled=pg >= n_pages - 1, on_click=_goto,
                      args=(pkey, labels[min(pg + 1, n_pages - 1)]))
        shown = pool[pg * PAGE:(pg + 1) * PAGE]

        st.caption(
            f"{len(pool)} of {len(feats)} features match"
            + (f" · showing {pg * PAGE + 1}-{pg * PAGE + len(shown)}" if pool else "")
            + f" · {len(fvals)} value(s) filled so far. Pick a recent value or type a new one; leave blank to skip."
        )

        if feats:
            with st.container(height=480, border=True):
                if not shown:
                    st.info(f"No feature matches '{q}'. If it is really new, add it below "
                            "(similar names are checked for you)." if q else "Nothing to show.")
                cols = st.columns(2)
                for i, f in enumerate(shown):
                    key = f"fv_{f}_{ver}"
                    opts = list(data.recent.get(f, []))  # most recently used values first
                    stored = fvals.get(f)
                    if key not in st.session_state and stored:
                        if stored not in opts:
                            opts = opts + [stored]
                        st.session_state[key] = stored  # restore after being off-screen
                    with cols[i % 2]:
                        st.selectbox(
                            ("🆕 " + f) if f in new_from_others else f,
                            opts,
                            index=None,
                            accept_new_options=True,
                            placeholder="select or type...",
                            key=key,
                        )
        else:
            st.info("No features yet. Add the first ones below.")

        # ---------------------------------------------------- add a brand-new feature
        st.markdown("**➕ Add a new feature**")
        nf = st.session_state.setdefault("nf_ver", 0)
        pend = st.session_state.setdefault(f"pending_{ver}", {})
        name = st.selectbox(
            "New feature name",
            sorted(feats, key=str.lower),
            index=None,
            accept_new_options=True,
            placeholder="Type a name: existing similar features show up as you type",
            key=f"nf_name_{ver}_{nf}",
        )
        name = (name or "").strip()
        similar, exact = similar_features(name, feats)
        can_add = bool(name)
        if name and exact:
            st.warning(f"**{exact}** already exists, so use that one instead of adding a new feature.")
            st.button(f"🔍 Find '{exact}' in the list above", key=f"jump_{ver}_{nf}",
                      on_click=_set_search, args=(exact, ver))
            can_add = False
        elif name and similar:
            st.warning("Similar features already exist: " + ", ".join(f"**{x}**" for x in similar)
                       + ". If one of them means the same thing, use it instead.")
            for x in similar[:3]:
                st.button(f"🔍 Find '{x}' in the list above", key=f"jump_{x}_{ver}_{nf}",
                          on_click=_set_search, args=(x, ver))
            can_add = st.checkbox("This is genuinely a different feature", key=f"nf_ok_{ver}_{nf}")
        elif name:
            st.success("No similar feature exists yet, so it is safe to add.")
        nval = st.text_input("Value for this product", key=f"nf_val_{ver}_{nf}")
        if st.button("Add this feature to the product", disabled=not can_add, key=f"nf_add_{ver}_{nf}"):
            if not nval.strip():
                st.error("Enter a value for the new feature.")
            elif norm(name) in RESERVED:
                st.error(f"'{name}' is a reserved column name; choose another name.")
            else:
                pend[name] = nval.strip()
                st.session_state["nf_ver"] = nf + 1  # fresh, empty inputs
                rerun_form()
        if pend:
            st.markdown("**New features for this product:**")
            for pn, pv in list(pend.items()):
                p1, p2 = st.columns([6, 1])
                p1.write(f"• **{pn}** = {pv}")
                if p2.button("✕", key=f"rm_{ver}_{pn}"):
                    pend.pop(pn)
                    rerun_form()

    st.divider()
    if st.button("💾 Save product", type="primary", width="stretch", disabled=not me):
        problems = []
        if not (bank and category and acc_type and product):
            problems.append("Bank, Bank Category, Account Type and Product Name are all required.")
        if problems:
            for p in problems:
                st.error(p)
        else:
            row = {
                "Bank Category": category,
                "Bank": bank,
                "Account Type": acc_type,
                "Product Name": product,
                "Added By": me,
                "Added At": datetime.now().strftime("%Y-%m-%d %H:%M"),
            }
            row.update(fvals)  # every value entered, including features that are off-screen right now
            new_feats = 0
            existing = {norm(x): x for x in feats}
            for f, v in st.session_state.get(f"pending_{ver}", {}).items():
                f = existing.get(norm(f), f)  # reuse the existing spelling if a teammate just added it
                if norm(f) in RESERVED:
                    st.error(f"'{f}' is a reserved column name; remove it and choose another name.")
                    st.stop()
                new_feats += f not in feats
                row[f] = v
            try:
                with st.spinner("Saving to the shared sheet..."):
                    row_no = save_row(row)
            except Exception as e:
                log.exception("Could not save product")
                code, msg, hint = explain_error(e)
                st.error(
                    "Could not save, so nothing was added. Your entries are still on the screen. " + hint
                )
                with st.expander("Technical details (send these if it keeps happening)"):
                    st.code(f"{type(e).__name__}\nGoogle status: {code}\nGoogle message: {msg}")
                st.stop()
            st.session_state["flash"] = (
                f"Saved **{product}** ({bank}) with {len(row) - len(FIXED)} feature values"
                + (f", including {new_feats} new feature(s)" if new_feats else "")
                + (f". It is in row {row_no} of the Google Sheet." if row_no else ".")
            )
            st.session_state["form_ver"] += 1  # fresh form; bank list now shows this bank first
            st.session_state["fvals"] = {}
            st.session_state["nf_ver"] = 0
            st.rerun()  # whole page: counts, tables and dropdowns pick up the new product


# ------------------------------------------------------------------ MY WORK
@st.fragment
def my_work():
    if not me:
        st.info("Pick your name in the sidebar to see your own products.")
    elif mine.empty:
        st.info("You haven't added anything yet. Use the **Add product** tab.")
    else:
        st.subheader(f"{me}'s products")
        filled = mine[feats].ne("").sum(axis=1) if feats else pd.Series(0, index=mine.index)
        m1, m2, m3 = st.columns(3)
        m1.metric("My products", len(mine))
        m2.metric("My banks", mine["Bank"].nunique())
        m3.metric("Avg features filled", f"{filled.mean():.1f} / {len(feats)}")
        chosen = feature_column_picker("mine")
        mv = mine[CORE + chosen + META].copy()
        mv.insert(4, "Features filled", filled.values)
        st.dataframe(mv, width="stretch", height=420)
        excel_download(mine, "my products (Excel)", f"{me}_products", "mine")
        st.caption("Low 'Features filled' on a product usually means values are still missing. "
                   "Fix them directly in the Google Sheet.")


# ------------------------------------------------------------------ BROWSE
@st.fragment
def browse():
    st.subheader("All products")
    f1, f2, f3, f4 = st.columns(4)
    sel_cat = f1.multiselect("Bank Category", sorted(df["Bank Category"].unique()))
    sel_bank = f2.multiselect("Bank", sorted(df["Bank"].unique()))
    sel_acc = f3.multiselect("Account Type", sorted(df["Account Type"].unique()))
    sel_by = f4.multiselect("Added By", sorted(df["Added By"].unique()))
    view = df
    if sel_by:
        view = view[view["Added By"].isin(sel_by)]
    if sel_cat:
        view = view[view["Bank Category"].isin(sel_cat)]
    if sel_bank:
        view = view[view["Bank"].isin(sel_bank)]
    if sel_acc:
        view = view[view["Account Type"].isin(sel_acc)]
    chosen = feature_column_picker("all")
    st.dataframe(view[CORE + chosen + META], width="stretch", height=460)
    st.caption(f"Showing {len(view)} of {len(df)} products and {len(chosen)} of {len(feats)} feature columns.")

    d1, d2 = st.columns([3, 1])
    with d1:
        excel_download(df, "Excel (all data, all columns)", "bank_products", "all")
    if d2.button("🔄 Refresh"):
        load_data.clear()
        st.rerun()
    st.caption("To fix or delete a row, edit it directly in the Google Sheet. The app picks changes up within seconds.")


# ------------------------------------------------------------------ OVERVIEW
@st.fragment
def overview():
    o1, o2 = st.columns(2)
    with o1:
        st.subheader("Products per bank")
        if len(df):
            st.bar_chart(df[df["Bank"] != ""].groupby("Bank").size().sort_values(ascending=False))
    with o2:
        st.subheader("Products per category")
        if len(df):
            st.bar_chart(df[df["Bank Category"] != ""].groupby("Bank Category").size())

    st.subheader(f"Feature coverage ({len(feats)} features)")
    if feats:
        cov = pd.DataFrame({
            "Feature": feats,
            "Products with value": [data.coverage[f] for f in feats],
            "Coverage %": [round(100 * data.coverage[f] / max(len(df), 1), 1) for f in feats],
        }).sort_values("Products with value", ascending=False)
        st.dataframe(cov, width="stretch", hide_index=True, height=420)
    else:
        st.info("No features yet — add your first product.")

    st.subheader("Team progress")
    if len(df):
        st.bar_chart(df[df["Added By"] != ""].groupby("Added By").size())

    st.subheader("Contributions")
    if len(df):
        st.dataframe(
            df.groupby(["Added By", "Bank"]).size().rename("Products").reset_index(),
            width="stretch", hide_index=True,
        )


with tab_add:
    add_form()
with tab_mine:
    my_work()
with tab_browse:
    browse()
with tab_overview:
    overview()
